"""Прокурорский шаблон из готовых находок; без повторной правовой экспертизы.

Блоки утверждения и подписи восстановлены из шаблона до 55f9236b.
"""

import re
from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt

from src.agents.enums.common.review import CheckStatus, RiskKind
from src.agents.schemas.common.review import Review, Risk
from src.agents.schemas.competence.sources import CompetenceRefs
from src.agents.schemas.parse.parsed_document import ParsedDocument
from src.agents.steps_tools.authority.requisites import parse_input
from src.agents.steps_tools.repeal.clauses import collect_repeal_targets
from src.agents.steps_tools.report.assemble import REPORT_ORDER
from src.common.config.config import settings

_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_OFFSETS = re.compile(r"\s*\(символы \d+\s*[–—-]\s*\d+\)")


def _xml_text(text: str) -> str:
    """Нормализация только вывода: исходные доказательства не меняются."""
    return _INVALID_XML.sub("\ufffd", text.replace("\f", "\n"))


_TEXT_WIDTH_CM = 16.5


_APPROVAL_INDENT_CM = 8.0


_APPROVAL_TITLE_LINES: tuple[str, ...] = (
    "Старший помощник прокурора",
    "Ульяновской области по взаимодействию",
    "с законодательным и исполнительными",
    "органами, органами местного",
    "самоуправления",
)


_APPROVAL_RANK = "старший советник юстиции"


_SIGN_TITLE_LINES: tuple[str, ...] = (
    "Прокурор отдела управления",
    "по надзору за исполнением",
    "федерального законодательства",
    "прокуратуры Ульяновской области",
)


_GENITIVE: dict[str, str] = {
    "приказ": "приказа",
    "постановление": "постановления",
    "распоряжение": "распоряжения",
    "закон": "закона",
    "кодекс": "кодекса",
    "положение": "положения",
    "регламент": "регламента",
    "административный регламент": "административного регламента",
    "указ": "указа",
    "решение": "решения",
}


_PLACEHOLDER = "_______________________________"


def _add_line(
    document: DocxDocument,
    text: str,
    *,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    bold: bool = False,
    indent: float = 0.0,
    left_indent: float = 0.0,
    size: int | None = None,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Cm(indent)
    if left_indent:
        paragraph.paragraph_format.left_indent = Cm(left_indent)
    run = paragraph.add_run(_xml_text(text))
    run.bold = bold
    if size is not None:
        run.font.size = Pt(size)


def _add_two_col_line(document: DocxDocument, left_text: str, right_text: str, *, size: int | None = None) -> None:
    """Строка с текстом слева и справа на одном уровне (чин … Ф.И.О.)."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(_TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    run = paragraph.add_run(_xml_text(f"{left_text}\t{right_text}"))
    if size is not None:
        run.font.size = Pt(size)


def _add_body_paragraph(document: DocxDocument, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(_xml_text(text))


def _add_blank(document: DocxDocument) -> None:
    document.add_paragraph()


def _genitive(doc_type: str | None) -> str:
    if not doc_type:
        return "нормативного правового акта"
    key = doc_type.strip().lower()
    return _GENITIVE.get(key, doc_type.strip())


def _approval_block(document: DocxDocument) -> None:
    left = WD_ALIGN_PARAGRAPH.LEFT
    ind = _APPROVAL_INDENT_CM
    _add_line(document, "УТВЕРЖДАЮ", align=left, left_indent=ind)
    for line in _APPROVAL_TITLE_LINES:
        _add_line(document, line, align=left, left_indent=ind)
    _add_blank(document)
    _add_line(document, _APPROVAL_RANK, align=left, left_indent=ind)
    _add_blank(document)
    _add_line(document, _PLACEHOLDER, align=left, left_indent=ind)
    _add_line(document, "(Ф.И.О. утверждающего лица)", align=left, left_indent=ind, size=10)
    _add_line(document, "__.__.____", align=left, left_indent=ind)


def _quoted(name: str) -> str:
    """Наименование в кавычках-ёлочках (если их ещё нет) — как в шапке заключения."""
    name = (name or "").strip()
    if not name:
        return ""
    return name if name.startswith("«") and name.endswith("»") else f"«{name}»"


def _signature_block(document: DocxDocument) -> None:
    left = WD_ALIGN_PARAGRAPH.LEFT
    _add_blank(document)
    for line in _SIGN_TITLE_LINES:
        _add_line(document, line, align=left)
    _add_blank(document)
    _add_two_col_line(document, _PLACEHOLDER, _PLACEHOLDER)
    _add_two_col_line(document, "(классный чин)", "(Ф.И.О.)", size=10)
    _add_blank(document)
    _add_line(document, "__.__.____", align=left)
    _add_blank(document)
    _add_line(document, "(Ф.И.О. исполнителя, тел.)", align=left, size=10)


def _location(text: str) -> str:
    """Смещения и дословные повторы адреса убираются только при печати."""
    return ", ".join(dict.fromkeys(part.strip() for part in _OFFSETS.sub("", text).split(",") if part.strip()))


def _finding_text(risk: Risk, index: int) -> str:
    connector = "Так, " if index == 0 else "Кроме того, " if index == 1 else "Также, "
    project, *others = risk.citations
    address = f" ({_location(project.location)})" if project.location else ""
    parts = [f"{connector}в проекте{address} предусмотрено: «{project.quote}».", risk.explanation]
    for citation in others:
        label = _location(citation.location) or "Цитируемая норма"
        parts.append(f"{label}: «{citation.quote}».")
    parts.append(risk.recommendation)
    return " ".join(parts)


def _setup_styles(document: DocxDocument) -> None:
    cfg = settings["agents.report.docx"]
    style = document.styles["Normal"]
    style.font.name = cfg.font_name
    style.font.size = Pt(float(cfg.font_size_pt))
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.first_line_indent = Cm(1.25)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.widow_control = True
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = Cm(float(cfg.left_margin_cm))
        section.right_margin = Cm(float(cfg.right_margin_cm))
        section.top_margin = section.bottom_margin = Cm(float(cfg.vertical_margin_cm))


def _subject(refs: CompetenceRefs | None, parsed: ParsedDocument | None) -> str:
    title = " ".join(refs.project_title.split()) if refs else ""
    # «В соответствии» внутри названия услуги ранее обрывало C1-заголовок.
    # Дополняем только однозначный дословный префикс до его парной кавычки.
    depth = title.count("«") - title.count("»")
    if depth > 0 and parsed is not None:
        header = " ".join(parse_input(parsed).split())
        if header.count(title) == 1:
            tail = header[header.index(title) + len(title) :]
            for index, char in enumerate(tail):
                depth += (char == "«") - (char == "»")
                if depth == 0:
                    return title + tail[: index + 1]
    return title


def _title_block(
    document: DocxDocument,
    refs: CompetenceRefs | None,
    source_name: str,
    parsed: ParsedDocument | None,
) -> None:
    form = _genitive(refs.form if refs else "")
    issuer = refs.issuer if refs else ""
    head, _, tail = issuer.partition(" ")
    genitives = {
        "министерство": "Министерства",
        "департамент": "Департамента",
        "управление": "Управления",
        "правительство": "Правительства",
        "администрация": "Администрации",
        "комитет": "Комитета",
    }
    issuer = " ".join(filter(None, (genitives.get(head.casefold(), head), tail)))
    subject = _quoted(_subject(refs, parsed)) or "«(наименование проекта)»"
    _add_blank(document)
    _add_line(document, "ЗАКЛЮЧЕНИЕ", bold=True)
    _add_line(document, f"на проект {form}")
    _add_line(document, issuer or _PLACEHOLDER)
    _add_line(document, subject)
    _add_blank(document)
    _add_body_paragraph(
        document, f"Прокуратурой области изучен проект {form} {issuer or _PLACEHOLDER} {subject} (далее – проект)."
    )
    if not refs or not refs.project_title:
        _add_body_paragraph(document, f"Представленный документ: {source_name}.")


def _project_context(document: DocxDocument, parsed: ParsedDocument | None) -> None:
    if parsed is None:
        return
    if parsed.act_basis and parsed.act_basis.preamble_text:
        preamble = parsed.act_basis.preamble_text
        markers = list(re.finditer(r"(?im)^\s*(?:В соответствии|Во исполнение|В целях|Руководствуясь)\b", preamble))
        if markers:
            preamble = " ".join(preamble[markers[-1].start() :].split())
            _add_body_paragraph(document, f"В качестве основания разработки проекта указано: «{preamble}».")
    targets = collect_repeal_targets(parsed.tree)
    if targets:
        _add_body_paragraph(document, "Проектом признаются утратившими силу (отменяются):")
        for target, _, _ in targets:
            unit = f"положение {target.subpoint}: " if target.is_subpoint and target.subpoint else ""
            _add_body_paragraph(document, unit + target.raw + ";")


def _complete(review: Review) -> bool:
    return {check.kind for check in review.checks} == set(REPORT_ORDER) and all(
        check.status in (CheckStatus.COMPLETED, CheckStatus.NOT_APPLICABLE) for check in review.checks
    )


def _resolution_text(review: Review) -> str:
    if review.findings:
        start = "Таким образом" if len(review.findings) == 1 else "В этой связи"
        return f"{start}, проект в указанной редакции принят быть не может и нуждается в доработке."
    if _complete(review):
        return "Таким образом, в пределах выполненных проверок замечания к проекту отсутствуют."
    return "Проверка проекта проведена не в полном объёме. Вывод о возможности его принятия не сформирован."


def _leadin_text(review: Review) -> str:
    kinds = {risk.kind for risk in review.findings}
    if RiskKind.COMPETENCE in kinds:
        return "Установлено, что проект содержит положения, выходящие за пределы полномочий органа."
    if RiskKind.COMPLIANCE in kinds:
        internal = "содержит внутренние противоречия, а также " if RiskKind.INTERNAL in kinds else ""
        return f"Установлено, что проект {internal}не соответствует актам более высокой юридической силы."
    if RiskKind.INTERNAL in kinds:
        return "Установлено, что проект содержит внутренние противоречия."
    if RiskKind.REPEAL in kinds:
        return "Установлено, что предусмотренный проектом перечень отменяемых актов и положений требует уточнения."
    return (
        "В пределах выполненных проверок замечаний к проекту не выявлено."
        if _complete(review)
        else "В рамках выполненных проверок подтверждённых замечаний не выявлено."
    )


def report_docx(
    review: Review,
    summary: str,
    *,
    source_name: str,
    analysis_date: str,
    competence_refs: CompetenceRefs | None = None,
    parsed_document: ParsedDocument | None = None,
) -> bytes:
    """Подписной шаблон; полные статусы и ограничения — в справке после подписи."""
    document = Document()
    _setup_styles(document)
    _approval_block(document)
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.keep_with_next = True
    document.paragraphs[-1].paragraph_format.keep_with_next = False
    _title_block(document, competence_refs, source_name, parsed_document)
    _project_context(document, parsed_document)
    _add_body_paragraph(document, _leadin_text(review))
    for index, risk in enumerate(review.findings):
        _add_body_paragraph(document, _finding_text(risk, index))
    _add_body_paragraph(document, _resolution_text(review))
    if not _complete(review) and review.findings:
        _add_body_paragraph(
            document, "Экспертиза проведена не в полном объёме; пределы проверки приведены в приложенной справке."
        )
    start_signature = len(document.paragraphs)
    _signature_block(document)
    for paragraph in document.paragraphs[start_signature:-1]:
        paragraph.paragraph_format.keep_with_next = True
    document.add_page_break()
    _add_line(document, "Справка об объёме проведённой проверки", bold=True)
    _add_line(document, f"Документ: {source_name}")
    _add_line(document, f"Дата экспертизы: {analysis_date}")
    for text in _xml_text(summary).splitlines():
        _add_body_paragraph(document, text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
