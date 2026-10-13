from io import BytesIO

import pytest
from docx import Document

from src.agents.enums.common.review import CheckStatus, RiskKind
from src.agents.enums.competence.review import CompetenceScope
from src.agents.schemas.common.review import Check, Citation, Review, Risk
from src.agents.schemas.competence.review import CompetenceReview
from src.agents.schemas.competence.sources import CompetenceRefs
from src.agents.schemas.parse.parsed_document import ParsedDocument
from src.agents.steps import build_report_node as report_module
from src.agents.steps_tools.document.act_basis import extract_act_basis
from src.agents.steps_tools.document.structure import build_document_structure_tree
from src.agents.steps_tools.report.assemble import REPORT_ORDER, body_paragraphs
from src.agents.steps_tools.report.render import report_docx
from src.common.config.config import settings


def finding(kind: RiskKind) -> Risk:
    return Risk(
        kind=kind,
        explanation="Нарушение установленного срока.",
        recommendation="Изменить срок.",
        citations=[Citation(source="project", location="пункт 1", quote="Срок 10 дней.")],
    )


def branch(kind: RiskKind, *, status: CheckStatus = CheckStatus.COMPLETED, findings=()) -> Review:
    return Review(
        findings=list(findings),
        checks=[
            Check(kind=kind, status=status, reason="Часть источников не прочитана." if status is CheckStatus.LIMITED else "")
        ],
    )


def test_docx_preserves_evidence_across_page_breaks_and_invalid_xml_characters():
    quote = "Первая страница\fВторая страница\tСрок 5 дней."
    risk = Risk(
        kind=RiskKind.COMPLIANCE,
        citations=[Citation(source="project", quote=quote)],
        explanation="Данные\x00источника\uffff",
        recommendation="Обычный текст: <акт> & право 😀",
    )
    document = Document(
        BytesIO(
            report_docx(
                Review(findings=[risk], checks=[]),
                "Проверка\x0bограничена",
                source_name="Проект\x01\ud800",
                analysis_date="2026-10-12",
            )
        )
    )
    text = "\n".join(p.text for p in document.paragraphs)
    for expected in (
        "Проект\ufffd",
        "Первая страница\nВторая страница\tСрок 5 дней.",
        "Данные\ufffdисточника\ufffd",
        "Обычный текст: <акт> & право 😀",
        "Проверка\ufffdограничена",
    ):
        assert expected in text
    assert risk.citations[0].quote == quote


@pytest.fixture
def events(monkeypatch):
    emitted = []
    monkeypatch.setattr(report_module, "emit_ui_event", lambda *args, **kwargs: emitted.append(kwargs))
    return emitted


async def test_report_preserves_branch_findings_statuses_and_docx_text(events):
    risks = [finding(kind) for kind in REPORT_ORDER]
    competence = CompetenceReview(
        **branch(RiskKind.COMPETENCE, findings=[risks[0]]).model_dump(), scope=CompetenceScope.PROVISIONS
    )
    internal = branch(RiskKind.INTERNAL, status=CheckStatus.LIMITED, findings=[risks[1]])
    legal = branch(RiskKind.COMPLIANCE, findings=[risks[2]])
    repeal = branch(RiskKind.REPEAL, findings=[risks[3]])
    result = await report_module.build_report_node(
        {
            "competence_review": competence,
            "internal_review": internal,
            "legal_review": legal,
            "repeal_review": repeal,
            "analysis_date": "2026-10-11",
        }
    )
    review, summary = result["review_result"], result["review_summary"]
    assert review.findings == risks
    assert review.checks == [competence.checks[0], internal.checks[0], legal.checks[0], repeal.checks[0]]
    assert settings["agents.report.incomplete"] in summary
    assert "Часть источников не прочитана." in summary
    body = body_paragraphs(review)
    assert events == [{"status": "success", "content": "\n\n".join(body), "findings_shown": 4, "outcome": "warn"}]
    document = Document(BytesIO(report_docx(review, summary, source_name="Проект", analysis_date="2026-10-11")))
    text = "\n".join(p.text for p in document.paragraphs)
    for risk in risks:
        assert risk.explanation in text and risk.recommendation in text
        assert all(c.quote in text for c in risk.citations)
    assert text.endswith(summary)


@pytest.mark.parametrize(
    "status,has_findings",
    [(CheckStatus.COMPLETED, False), (CheckStatus.LIMITED, False), (CheckStatus.NOT_CHECKED, False), (CheckStatus.LIMITED, True)],
)
def test_prosecutor_template_preserves_evidence_and_qualifies_empty_results(status, has_findings):
    risk = finding(RiskKind.INTERNAL)
    risk.citations[0].location = "п. 2, п. 2 (символы 5–20)"
    risk.citations.append(Citation(source="project", location="пункт 3", quote="Срок 5 дней."))
    review = Review(
        findings=[risk] if has_findings else [],
        checks=[Check(kind=k, status=status, reason="LAW:123 — не получен текст.") for k in REPORT_ORDER],
    )
    before = review.model_dump_json()
    refs = CompetenceRefs(
        issuer="Министерство социального развития Ульяновской области", form="приказ", project_title="Об утверждении Порядка"
    )
    doc = Document(
        BytesIO(
            report_docx(
                review,
                "Статус проверок:\nLAW:123 — не получен текст.",
                source_name="Проект.docx",
                analysis_date="2026-10-12",
                competence_refs=refs,
            )
        )
    )
    text = "\n".join(p.text for p in doc.paragraphs)
    main, appendix = text.split("Справка об объёме проведённой проверки")
    for expected in (
        "УТВЕРЖДАЮ",
        "ЗАКЛЮЧЕНИЕ",
        "на проект приказа",
        "Министерства социального развития",
        "«Об утверждении Порядка»",
        "Прокуратурой области изучен проект",
        "(классный чин)",
        "(Ф.И.О.)",
    ):
        assert expected in main
    assert "LAW:123" not in main and "LAW:123" in appendix
    assert "Статус проверок:" not in main and "символы" not in main
    if has_findings:
        assert "Так, в проекте (п. 2)" in main and "нуждается в доработке" in main
        assert all(c.quote in main for c in risk.citations)
        assert risk.explanation in main and risk.recommendation in main
    elif status is CheckStatus.COMPLETED:
        assert "в пределах выполненных проверок замечания к проекту отсутствуют" in main
    else:
        assert "Вывод о возможности его принятия не сформирован" in main
        assert "замечания к проекту отсутствуют" not in main
    assert review.model_dump_json() == before
    assert doc.styles["Normal"].font.size.pt == 14
    assert round(doc.sections[0].left_margin.cm, 1) == 2.5


def test_full_title_and_preamble_are_printed_without_repeating_the_cover():
    title = "Об утверждении Регламента услуги «Освобождение в соответствии с законодательством опекунов»"
    source = (
        "Министерство социального развития Ульяновской области\nПРИКАЗ\n"
        "Об утверждении Регламента услуги «Освобождение\n в соответствии с законодательством опекунов»\n"
        "В соответствии со статьёй 29 Федерального закона от 24.04.2008 № 48-ФЗ\nприказываю:\n"
        "1. Утвердить Регламент.\n2. Признать утратившим силу приказ от 01.01.2020 № 1-п."
    )
    parsed = ParsedDocument(enriched_text=source, tree=build_document_structure_tree(source), act_basis=extract_act_basis(source))
    refs = CompetenceRefs(
        form="приказ",
        issuer="Министерство социального развития Ульяновской области",
        project_title="Об утверждении Регламента услуги «Освобождение",
    )
    before = refs.model_dump_json()
    doc = Document(
        BytesIO(
            report_docx(
                Review(findings=[], checks=[]),
                "Проверки не завершены.",
                source_name="Проект.docx",
                analysis_date="2026-10-12",
                competence_refs=refs,
                parsed_document=parsed,
            )
        )
    )
    paragraphs = [p.text for p in doc.paragraphs]
    assert f"«{title}»" in paragraphs
    basis = next(p for p in paragraphs if p.startswith("В качестве основания"))
    assert "со статьёй 29" in basis and "Министерство" not in basis and "ПРИКАЗ" not in basis
    assert "Проектом признаются утратившими силу (отменяются):" in paragraphs
    assert any("01.01.2020" in p and "1-п" in p for p in paragraphs)
    assert refs.model_dump_json() == before


async def test_deduplication_preserves_distinct_places_sources_and_citations(events):
    original = finding(RiskKind.REPEAL)
    duplicate = original.model_copy(deep=True)
    duplicate.explanation = "Нарушение  установленного\nсрока."
    duplicate.recommendation = " Изменить   срок. "
    duplicate.citations[0].location = " пункт   1 "
    duplicate.citations[0].quote = "Срок  10\nдней."
    another_place = original.model_copy(deep=True)
    another_place.citations[0].location = "пункт 2"
    another_source = original.model_copy(deep=True)
    another_source.citations.append(Citation(source="LAW:1", location="статья 1", quote="Срок 5 дней."))
    distinct_source = another_source.model_copy(deep=True)
    distinct_source.citations[1].source = "LAW:2"
    distinct_quote = another_source.model_copy(deep=True)
    distinct_quote.citations[1].quote = "Срок 3 дня."
    risks = [original, another_place, another_source, distinct_source, distinct_quote]
    result = await report_module.build_report_node(
        {"repeal_review": branch(RiskKind.REPEAL, findings=[original, duplicate, *risks[1:], original])}
    )
    assert result["review_result"].findings == risks


async def test_missing_compliance_does_not_limit_repeal_or_lose_its_findings(events):
    risk = finding(RiskKind.REPEAL)
    repeal = branch(RiskKind.REPEAL, findings=[risk])
    result = await report_module.build_report_node({"repeal_review": repeal})
    review = result["review_result"]
    assert review.findings == [risk]
    assert [c.kind for c in review.checks] == list(REPORT_ORDER)
    assert [c.status for c in review.checks] == [CheckStatus.NOT_CHECKED] * 3 + [CheckStatus.COMPLETED]
    assert review.checks[-1] == repeal.checks[0]
    assert settings["agents.report.incomplete"] in result["review_summary"]


@pytest.mark.parametrize("state", [{}, {"internal_review": Review(findings=[], checks=[])}])
async def test_absent_results_and_absent_status_are_not_a_clean_review(state, events):
    result = await report_module.build_report_node(state)
    assert all(c.status is CheckStatus.NOT_CHECKED for c in result["review_result"].checks)
    assert settings["agents.report.not_checked"] in result["review_summary"]
    assert settings["agents.report.no_findings"] not in result["review_summary"]
    assert events[0]["outcome"] == "warn"


@pytest.mark.parametrize(
    ("scope", "status", "with_finding", "stopped"),
    [
        (CompetenceScope.ACT, CheckStatus.COMPLETED, True, True),
        (CompetenceScope.ACT, CheckStatus.LIMITED, True, False),
        (CompetenceScope.ACT, CheckStatus.COMPLETED, False, False),
        (CompetenceScope.PROVISIONS, CheckStatus.COMPLETED, True, False),
    ],
)
async def test_early_stop_requires_confirmed_whole_act_result(scope, status, with_finding, stopped, events):
    risks = [finding(RiskKind.COMPETENCE)] if with_finding else []
    competence = CompetenceReview(**branch(RiskKind.COMPETENCE, status=status, findings=risks).model_dump(), scope=scope)
    result = await report_module.build_report_node({"competence_review": competence})
    review = result["review_result"]
    assert review.findings == risks
    assert review.checks[0] == competence.checks[0]
    assert all(c.status is CheckStatus.NOT_CHECKED for c in review.checks[1:])
    assert ("завершена досрочно" in result["review_summary"]) is stopped
    assert all(("полномочий на принятие акта в целом" in c.reason) is stopped for c in review.checks[1:])


async def test_complete_empty_review_has_clear_outcome(events):
    result = await report_module.build_report_node(
        {
            "competence_review": CompetenceReview(**branch(RiskKind.COMPETENCE).model_dump(), scope=CompetenceScope.NONE),
            "internal_review": branch(RiskKind.INTERNAL),
            "legal_review": branch(RiskKind.COMPLIANCE),
            "repeal_review": branch(RiskKind.REPEAL),
        }
    )
    assert result["review_result"].findings == []
    assert settings["agents.report.no_findings"] in result["review_summary"]
    assert settings["agents.report.incomplete"] not in result["review_summary"]
    assert events[0]["outcome"] == "clear"
